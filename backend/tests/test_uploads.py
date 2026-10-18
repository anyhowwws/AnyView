"""Resume upload tickets and the upload → text path at interview start (ADR 0010)."""

import io

import pytest
from docx import Document
from fastapi.testclient import TestClient

from anyview.api import interview
from anyview.llm.fake import FakeModel
from anyview.main import app
from anyview.uploads.extract import DOCX, MAX_UPLOAD_BYTES, PDF, docx_to_text
from anyview.uploads.store import GcsUploadStore, UploadTooLarge

UID = "uid-jordan"  # matches conftest.FAKE_TOKENS["good"]


@pytest.fixture
def client(signed_in):
    fake = FakeModel()
    app.dependency_overrides[interview.get_model] = lambda: fake
    c = TestClient(app, headers={"Authorization": "Bearer good"})
    c.fake = fake  # type: ignore[attr-defined]
    yield c
    app.dependency_overrides.pop(interview.get_model, None)


def _ticket(client, content_type=PDF, size=1234, filename="cv.pdf"):
    return client.post(
        "/uploads", json={"filename": filename, "content_type": content_type, "size_bytes": size}
    )


def _docx_bytes(paragraphs, table=None) -> bytes:
    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    if table:
        t = doc.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, cell in enumerate(row):
                t.cell(r, c).text = cell
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# --- /uploads --------------------------------------------------------------


def test_ticket_names_object_under_the_users_prefix(client):
    r = _ticket(client)
    assert r.status_code == 200, r.text
    ticket = r.json()
    assert ticket["object_name"].startswith(f"uploads/{UID}/")
    assert ticket["object_name"].endswith(".pdf")
    # The in-memory store points the browser back at the API's fake-mode route.
    assert ticket["upload_url"] == f"/uploads/memory/{ticket['object_name']}"
    assert ticket["expires_in_seconds"] == 900


def test_docx_ticket_gets_docx_extension(client):
    assert _ticket(client, DOCX, filename="cv.docx").json()["object_name"].endswith(".docx")


def test_filename_never_reaches_the_object_name(client):
    name = _ticket(client, filename="../../etc/passwd.pdf").json()["object_name"]
    assert ".." not in name and "passwd" not in name


def test_wrong_type_is_415_and_too_big_is_413(client):
    assert _ticket(client, "image/png").status_code == 415
    assert _ticket(client, PDF, size=11 * 1024 * 1024).status_code == 413


def test_uploads_require_sign_in(signed_in):
    r = TestClient(app).post(
        "/uploads", json={"filename": "cv.pdf", "content_type": PDF, "size_bytes": 1}
    )
    assert r.status_code == 401


# --- extraction ------------------------------------------------------------


def test_docx_to_text_reads_paragraphs_and_tables():
    data = _docx_bytes(
        ["Jordan Lee", "", "", "Senior Cloud Architect"],
        table=[["Acme Corp", "2021–present"], ["Globex", "2018–2021"]],
    )
    text = docx_to_text(data)
    assert text.startswith("Jordan Lee")
    assert "\n\n\n" not in text  # blank runs collapsed
    assert "Acme Corp | 2021–present" in text


def test_docx_to_text_rejects_garbage():
    from anyview.uploads.extract import ExtractionError

    with pytest.raises(ExtractionError):
        docx_to_text(b"not a docx")


# --- start with an upload ----------------------------------------------------


def _setup(**overrides):
    base = {"company": "Initech", "role": "Cloud Architect", "seniority": "senior"}
    return {**base, **overrides}


def test_start_with_uploaded_docx_uses_its_text(client, store):
    name = _ticket(client, DOCX, filename="cv.docx").json()["object_name"]
    store.put(
        name, _docx_bytes(["Jordan Lee — Cloud Solutions Architect"] + ["Led migrations."] * 5)
    )
    r = client.post(
        "/interview/start", json={"session_id": "sess-0001", "setup": _setup(resume_object=name)}
    )
    assert r.status_code == 200, r.text
    session = r.json()
    assert session["setup"]["resume"].startswith("Jordan Lee")
    assert session["setup"]["resume_object"] == name
    assert [c["call"] for c in client.fake.calls] == ["plan"]  # no model call for DOCX


def test_start_with_uploaded_pdf_asks_the_model_to_transcribe(client, store):
    name = _ticket(client, PDF).json()["object_name"]
    store.put(name, b"%PDF-1.7 fake bytes")
    r = client.post(
        "/interview/start", json={"session_id": "sess-0001", "setup": _setup(resume_object=name)}
    )
    assert r.status_code == 200, r.text
    calls = client.fake.calls
    assert [c["call"] for c in calls] == ["extract", "plan"]
    assert calls[0]["attachment"] == (len(b"%PDF-1.7 fake bytes"), PDF)
    assert "Jordan Lee" in r.json()["setup"]["resume"]


def test_start_refuses_someone_elses_object(client, store):
    name = "uploads/uid-other/abc.pdf"
    store.put(name, b"%PDF")
    r = client.post(
        "/interview/start", json={"session_id": "sess-0001", "setup": _setup(resume_object=name)}
    )
    assert r.status_code == 403


def test_start_with_missing_object_is_404(client):
    r = client.post(
        "/interview/start",
        json={"session_id": "sess-0001", "setup": _setup(resume_object=f"uploads/{UID}/gone.pdf")},
    )
    assert r.status_code == 404


def test_start_takes_an_upload_or_text_but_not_both(client):
    assert (
        client.post("/interview/start", json={"session_id": "s", "setup": _setup()}).status_code
        == 422
    )
    both = _setup(resume="x" * 60, resume_object=f"uploads/{UID}/a.pdf")
    assert (
        client.post("/interview/start", json={"session_id": "s", "setup": both}).status_code == 422
    )


def test_empty_docx_is_a_readable_422(client, store):
    name = _ticket(client, DOCX, filename="cv.docx").json()["object_name"]
    store.put(name, _docx_bytes(["hi"]))
    r = client.post(
        "/interview/start", json={"session_id": "sess-0001", "setup": _setup(resume_object=name)}
    )
    assert r.status_code == 422
    assert "at least 50 characters" in r.json()["detail"]


def test_a_recording_is_transcribed_on_the_fast_model():
    # ADR 0039: Pro took 4-50 s for a few words; transcription is mechanical.
    from anyview.uploads.extract import audio_to_text

    fake = FakeModel({"transcribe": "I led the migration."})
    assert audio_to_text(b"RIFF", "audio/wav", model=fake, session_id="sess-0001") == (
        "I led the migration."
    )
    assert fake.calls[0]["call"] == "transcribe" and fake.calls[0]["fast"] is True


# --- The 10 MB limit holds after the ticket too ------------------------------


def test_ticket_tells_the_browser_which_headers_to_send(client):
    # The in-memory store signs nothing; the real one signs the size range.
    assert _ticket(client).json()["upload_headers"] == {}


def test_fake_mode_put_refuses_more_than_10_mb(client):
    name = _ticket(client).json()["object_name"]
    r = client.put(f"/uploads/memory/{name}", content=b"x" * (MAX_UPLOAD_BYTES + 1))
    assert r.status_code == 413


def test_an_oversize_object_is_never_read(client, store):
    # A browser that lied about size_bytes, against a URL signed before the
    # size header existed: refused when read, before any model call.
    name = f"uploads/{UID}/big.pdf"
    store.put(name, b"%PDF" + b"x" * MAX_UPLOAD_BYTES)
    r = client.post(
        "/interview/start",
        json={
            "session_id": "sess-0001",
            "setup": {
                "company": "Initech",
                "role": "SRE",
                "seniority": "senior",
                "resume_object": name,
            },
        },
    )
    assert r.status_code == 413
    assert client.fake.calls == []


class _Blob:
    def __init__(self, size):
        self.size = size

    def download_as_bytes(self):
        return b"data"


class _Bucket:
    def __init__(self, blobs):
        self.blobs = blobs
        self.signed = None

    def get_blob(self, name):
        return self.blobs.get(name)

    def blob(self, name):
        bucket = self

        class _Signer:
            def generate_signed_url(self, **kw):
                bucket.signed = kw
                return "https://storage.example/signed"

        return _Signer()


def _gcs(blobs):
    gcs = GcsUploadStore.__new__(GcsUploadStore)  # skip __init__: no Google client
    gcs._bucket = _Bucket(blobs)
    gcs._signer = "anyview-api@example.iam.gserviceaccount.com"
    gcs._access_token = lambda: "token"
    return gcs


def test_gcs_signs_the_size_range_into_the_url():
    gcs = _gcs({})
    gcs.signed_put_url("uploads/u/x.pdf", PDF)
    assert gcs._bucket.signed["headers"] == {"x-goog-content-length-range": f"0,{MAX_UPLOAD_BYTES}"}
    assert gcs.put_headers() == gcs._bucket.signed["headers"]


def test_gcs_read_checks_the_size_before_downloading():
    gcs = _gcs({"ok": _Blob(10), "big": _Blob(MAX_UPLOAD_BYTES + 1)})
    assert gcs.read("ok") == b"data"
    assert gcs.read("missing") is None
    with pytest.raises(UploadTooLarge):
        gcs.read("big")
